"""
Gera arquivos Excel para compartilhamento dos dados do dashboard.

O módulo não depende do Streamlit: recebe apenas o DataFrame já carregado e
filtrado. Assim, a exportação não dispara consultas adicionais ao Supabase e
continua simples de testar e reaproveitar em outros pontos da aplicação.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from numbers import Number
from typing import Final

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from app_postos.core.config import Columns, RawColumns


class DataExportError(Exception):
    """Erro de domínio para uma exportação que não pôde ser gerada."""


# Apenas campos persistidos são exportados. Colunas derivadas para a interface
# (como ``ano_mes`` e ``oficina_mp``) não devem virar um novo contrato externo.
EXPORT_COLUMNS: Final[tuple[tuple[str, str], ...]] = (
    (Columns.FRETE, RawColumns.FRETE),
    (Columns.MP, RawColumns.MP),
    (Columns.OFICINA, RawColumns.OFICINA),
    (Columns.DATA_EFETIVOS, RawColumns.DATA_EFETIVOS),
    (Columns.QTD_EFETIVOS, RawColumns.QTD_EFETIVOS),
    (Columns.DATA_TRABALHADOS, RawColumns.DATA_TRABALHADOS),
    (Columns.QTD_TRABALHADOS, RawColumns.QTD_TRABALHADOS),
    (Columns.CONTRATACOES, RawColumns.CONTRATACAO),
    (Columns.DEMISSOES, RawColumns.DEMISSAO),
    (Columns.SEMANA, RawColumns.SEMANA),
)

_DATE_COLUMNS: Final[frozenset[str]] = frozenset(
    {Columns.DATA_EFETIVOS, Columns.DATA_TRABALHADOS}
)
_TEXT_COLUMNS: Final[frozenset[str]] = frozenset(
    {Columns.FRETE, Columns.MP, Columns.OFICINA}
)
_FORMULA_PREFIXES: Final[tuple[str, ...]] = ("=", "+", "-", "@")
_HEADER_FILL: Final[str] = "18C99E"
_HEADER_FONT: Final[str] = "FFFFFF"
_MAX_COLUMN_WIDTH: Final[int] = 42


def build_excel_export(df: pd.DataFrame, scope_description: str) -> bytes:
    """
    Retorna uma planilha .xlsx pronta para download com o DataFrame informado.

    ``scope_description`` é exibido na aba de informações para deixar claro se
    o arquivo representa os filtros ativos ou a base completa. O DataFrame não
    é alterado durante o processo.
    """
    if df.empty:
        raise DataExportError("Não há registros para exportar.")

    _validate_export_schema(df)
    data = _prepare_export_data(df)

    try:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Dados"
        _write_data_sheet(worksheet, data)
        _write_info_sheet(workbook, scope_description, len(data))

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    except DataExportError:
        raise
    except Exception as exc:  # noqa: BLE001 - converte erro técnico no contrato do serviço.
        raise DataExportError(f"Não foi possível montar o arquivo Excel: {exc}") from exc


def _validate_export_schema(df: pd.DataFrame) -> None:
    missing_columns = [column for column, _ in EXPORT_COLUMNS if column not in df.columns]
    if missing_columns:
        raise DataExportError(
            "Os dados não possuem todas as colunas necessárias para a exportação: "
            f"{', '.join(missing_columns)}."
        )


def _prepare_export_data(df: pd.DataFrame) -> pd.DataFrame:
    """Seleciona e ordena uma cópia independente dos campos exportáveis."""
    columns = [column for column, _ in EXPORT_COLUMNS]
    return (
        df.loc[:, columns]
        .copy()
        .sort_values(
            [Columns.DATA_EFETIVOS, Columns.SEMANA, Columns.OFICINA, Columns.MP],
            kind="stable",
        )
        .reset_index(drop=True)
    )


def _write_data_sheet(worksheet, data: pd.DataFrame) -> None:
    headers = [label for _, label in EXPORT_COLUMNS]
    worksheet.append(headers)

    for row in data.itertuples(index=False, name=None):
        worksheet.append(
            [_to_excel_value(value, column) for value, (column, _) in zip(row, EXPORT_COLUMNS)]
        )

    header_fill = PatternFill(fill_type="solid", fgColor=_HEADER_FILL)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color=_HEADER_FONT, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 22

    table = Table(displayName="DadosPostos", ref=worksheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    worksheet.add_table(table)

    _format_columns(worksheet)
    _fit_column_widths(worksheet)


def _to_excel_value(value: object, column: str) -> object:
    """Converte nulos/tipos pandas e evita que textos sejam fórmulas no Excel."""
    if value is None or pd.isna(value):
        return None
    if column in _DATE_COLUMNS:
        return pd.Timestamp(value).to_pydatetime()
    if column in _TEXT_COLUMNS:
        text = str(value)
        return f"'{text}" if text.startswith(_FORMULA_PREFIXES) else text
    if isinstance(value, Number) and hasattr(value, "item"):
        return value.item()
    return value


def _format_columns(worksheet) -> None:
    column_indexes = {column: index + 1 for index, (column, _) in enumerate(EXPORT_COLUMNS)}
    for column in _DATE_COLUMNS:
        for cell in worksheet.iter_cols(
            min_col=column_indexes[column], max_col=column_indexes[column], min_row=2
        ):
            for value in cell:
                value.number_format = "DD/MM/YYYY"

    for column, index in column_indexes.items():
        if column not in _TEXT_COLUMNS:
            worksheet.column_dimensions[worksheet.cell(1, index).column_letter].alignment = Alignment(
                horizontal="center"
            )


def _fit_column_widths(worksheet) -> None:
    """Limita a largura para manter o arquivo legível mesmo com textos longos."""
    for column_cells in worksheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(width, _MAX_COLUMN_WIDTH)


def _write_info_sheet(workbook: Workbook, scope_description: str, row_count: int) -> None:
    worksheet = workbook.create_sheet("Informações")
    worksheet.append(["Exportação de dados"])
    worksheet.append(["Escopo", scope_description])
    worksheet.append(["Registros exportados", row_count])
    worksheet.append(["Gerado em", datetime.now().astimezone().strftime("%d/%m/%Y %H:%M:%S %z")])
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 62
    worksheet["A1"].font = Font(bold=True, size=14, color=_HEADER_FONT)
    worksheet["A1"].fill = PatternFill(fill_type="solid", fgColor=_HEADER_FILL)
    worksheet.merge_cells("A1:B1")
