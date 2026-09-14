"""Gera relatórios consolidados por semana em Excel e PDF.

O módulo recebe somente o DataFrame limpo da aplicação e não depende do
Streamlit. Isso mantém a agregação e a geração dos arquivos testáveis sem
acesso ao Supabase ou à interface.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
import math
from numbers import Integral, Real
from typing import Final

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from app_postos.core.config import Columns
from app_postos.core.utils import safe_div
from app_postos.services.export_service import DataExportError


class WeeklyReportColumns:
    """Contrato interno do resumo semanal."""

    ANO = "ano"
    SEMANA = "semana"
    TOTAL_EFETIVO = "total_efetivo"
    TOTAL_TRABALHADO = "total_trabalhado"
    TOTAL_AUSENCIA = "total_ausencia"
    ABSENTEISMO = "absenteismo"


_REQUIRED_COLUMNS: Final[tuple[str, ...]] = (
    Columns.ANO,
    Columns.SEMANA,
    Columns.QTD_EFETIVOS,
    Columns.QTD_TRABALHADOS,
)
_REPORT_COLUMNS: Final[tuple[tuple[str, str], ...]] = (
    (WeeklyReportColumns.ANO, "Ano"),
    (WeeklyReportColumns.SEMANA, "Semana"),
    (WeeklyReportColumns.TOTAL_EFETIVO, "Total Efetivo"),
    (WeeklyReportColumns.TOTAL_TRABALHADO, "Total Trabalhado"),
    (WeeklyReportColumns.TOTAL_AUSENCIA, "Total Ausência"),
    (WeeklyReportColumns.ABSENTEISMO, "Absenteísmo"),
)
_HEADER_FILL: Final[str] = "0E2A3F"
_HEADER_FONT: Final[str] = "FFFFFF"


def build_weekly_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Consolida totais por ano e semana sem alterar o DataFrame recebido."""
    source = _prepare_source_data(df)

    grouped = (
        source.groupby([Columns.ANO, Columns.SEMANA], as_index=False, dropna=False)[
            [Columns.QTD_EFETIVOS, Columns.QTD_TRABALHADOS]
        ]
        .sum()
        .sort_values([Columns.ANO, Columns.SEMANA], kind="stable")
        .reset_index(drop=True)
        .rename(
            columns={
                Columns.ANO: WeeklyReportColumns.ANO,
                Columns.SEMANA: WeeklyReportColumns.SEMANA,
                Columns.QTD_EFETIVOS: WeeklyReportColumns.TOTAL_EFETIVO,
                Columns.QTD_TRABALHADOS: WeeklyReportColumns.TOTAL_TRABALHADO,
            }
        )
    )
    grouped[WeeklyReportColumns.TOTAL_AUSENCIA] = (
        grouped[WeeklyReportColumns.TOTAL_EFETIVO]
        - grouped[WeeklyReportColumns.TOTAL_TRABALHADO]
    )
    grouped[WeeklyReportColumns.ABSENTEISMO] = grouped.apply(
        lambda row: safe_div(
            row[WeeklyReportColumns.TOTAL_AUSENCIA],
            row[WeeklyReportColumns.TOTAL_EFETIVO],
        )
        * 100,
        axis=1,
    )
    return grouped[[column for column, _ in _REPORT_COLUMNS]]


def build_weekly_excel_report(df: pd.DataFrame, scope_description: str) -> bytes:
    """Retorna o resumo semanal em um arquivo XLSX formatado e auditável."""
    summary = build_weekly_summary(df)
    try:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Resumo Semanal"
        _write_excel_summary(worksheet, summary)
        _write_excel_info(workbook, scope_description, len(summary))
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    except DataExportError:
        raise
    except Exception as exc:  # noqa: BLE001 - contrato estável para a camada de UI.
        raise DataExportError(
            f"Não foi possível montar o resumo semanal em Excel: {exc}"
        ) from exc


def build_weekly_pdf_report(df: pd.DataFrame, scope_description: str) -> bytes:
    """Retorna o resumo semanal em PDF, com cabeçalho repetido entre páginas."""
    summary = build_weekly_summary(df)
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    except ImportError as exc:
        raise DataExportError(
            "A geração de PDF não está disponível porque a dependência reportlab não foi instalada."
        ) from exc

    try:
        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=14 * mm,
            leftMargin=14 * mm,
            topMargin=13 * mm,
            bottomMargin=13 * mm,
            title="Resumo semanal de postos de trabalho",
            author="Gestão de Postos de Trabalho",
            subject=str(scope_description),
        )

        # Strings simples permitem que o TEXTCOLOR da TableStyle controle o
        # cabeçalho. Paragraphs com BodyText manteriam a cor preta interna.
        table_data: list[list[object]] = [[label for _, label in _REPORT_COLUMNS]]
        for row in summary.itertuples(index=False):
            table_data.append(
                [
                    str(int(row.ano)),
                    str(int(row.semana)),
                    _format_integer(row.total_efetivo),
                    _format_integer(row.total_trabalhado),
                    _format_integer(row.total_ausencia),
                    _format_percentage(row.absenteismo),
                ]
            )

        table = Table(
            table_data,
            repeatRows=1,
            colWidths=[25 * mm, 25 * mm, 39 * mm, 43 * mm, 39 * mm, 35 * mm],
            hAlign="CENTER",
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_HEADER_FILL}")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 9),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("FONTSIZE", (0, 1), (-1, -1), 9),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{_HEADER_FILL}")),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )

        story = [table]
        document.build(story, onFirstPage=_draw_pdf_footer, onLaterPages=_draw_pdf_footer)
        return output.getvalue()
    except DataExportError:
        raise
    except Exception as exc:  # noqa: BLE001 - converte detalhes técnicos para o domínio.
        raise DataExportError(
            f"Não foi possível montar o resumo semanal em PDF: {exc}"
        ) from exc


def _prepare_source_data(df: pd.DataFrame) -> pd.DataFrame:
    if not isinstance(df, pd.DataFrame):
        raise DataExportError("Os dados informados para o relatório são inválidos.")
    if df.empty:
        raise DataExportError("Não há registros para exportar.")

    missing = [column for column in _REQUIRED_COLUMNS if column not in df.columns]
    if missing:
        raise DataExportError(
            "Os dados não possuem todas as colunas necessárias para o resumo semanal: "
            f"{', '.join(missing)}."
        )

    numeric_columns = list(_REQUIRED_COLUMNS)
    converted = {column: pd.to_numeric(df[column], errors="coerce") for column in numeric_columns}
    invalid = [column for column, values in converted.items() if values.isna().any()]
    if invalid:
        raise DataExportError(
            "Há valores vazios ou não numéricos nas colunas do resumo semanal: "
            f"{', '.join(invalid)}."
        )

    non_finite = [
        column
        for column, values in converted.items()
        if (~values.map(lambda value: math.isfinite(float(value)))).any()
    ]
    if non_finite:
        raise DataExportError(
            "Há valores infinitos nas colunas do resumo semanal: "
            f"{', '.join(non_finite)}."
        )

    if ((converted[Columns.SEMANA] % 1 != 0) | ~converted[Columns.SEMANA].between(1, 53)).any():
        raise DataExportError("A coluna semana deve conter números inteiros entre 1 e 53.")
    if ((converted[Columns.ANO] % 1 != 0) | ~converted[Columns.ANO].between(1, 9999)).any():
        raise DataExportError("A coluna ano deve conter anos válidos.")
    for column in (Columns.QTD_EFETIVOS, Columns.QTD_TRABALHADOS):
        if (converted[column] < 0).any():
            raise DataExportError(f"A coluna {column} não pode conter valores negativos.")
        if (converted[column] % 1 != 0).any():
            raise DataExportError(f"A coluna {column} deve conter quantidades inteiras.")

    normalized = pd.DataFrame(converted, index=df.index)
    for column in numeric_columns:
        normalized[column] = normalized[column].astype("int64")
    return normalized


def _write_excel_summary(worksheet, summary: pd.DataFrame) -> None:
    headers = [label for _, label in _REPORT_COLUMNS]
    worksheet.append(headers)

    for row_index, row in enumerate(summary.itertuples(index=False), start=2):
        worksheet.append(
            [
                int(row.ano),
                int(row.semana),
                _native_number(row.total_efetivo),
                _native_number(row.total_trabalhado),
                f"=C{row_index}-D{row_index}",
                f'=IF(C{row_index}=0,"",E{row_index}/C{row_index})',
            ]
        )

    header_fill = PatternFill(fill_type="solid", fgColor=_HEADER_FILL)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color=_HEADER_FONT, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    border = Border(
        left=Side(style="thin", color=_HEADER_FILL),
        right=Side(style="thin", color=_HEADER_FILL),
        top=Side(style="thin", color=_HEADER_FILL),
        bottom=Side(style="thin", color=_HEADER_FILL),
    )
    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = border
            if cell.row > 1:
                cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFF")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 24
    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 12
    for column in ("C", "D", "E"):
        worksheet.column_dimensions[column].width = 21
        for cell in worksheet[column][1:]:
            cell.number_format = "#,##0"
    worksheet.column_dimensions["F"].width = 18
    for cell in worksheet["F"][1:]:
        cell.number_format = "0.0%"
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    table = Table(displayName="ResumoSemanal", ref=worksheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _write_excel_info(workbook: Workbook, scope_description: str, row_count: int) -> None:
    worksheet = workbook.create_sheet("Informações")
    worksheet.append(["Relatório semanal"])
    worksheet.append(["Escopo", str(scope_description)])
    worksheet.append(["Semanas exportadas", int(row_count)])
    worksheet.append(["Gerado em", datetime.now().astimezone().strftime("%d/%m/%Y %H:%M:%S %z")])
    worksheet.append(
        [
            "Cálculo",
            "Ausência = Total Efetivo - Total Trabalhado; "
            "Absenteísmo = Ausência / Total Efetivo.",
        ]
    )
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 76
    worksheet["A1"].font = Font(bold=True, size=14, color=_HEADER_FONT)
    worksheet["A1"].fill = PatternFill(fill_type="solid", fgColor=_HEADER_FILL)
    worksheet.merge_cells("A1:B1")


def _native_number(value: object) -> int | float:
    if isinstance(value, Integral):
        return int(value)
    if isinstance(value, Real):
        numeric = float(value)
        return int(numeric) if numeric.is_integer() else numeric
    return float(value)


def _format_integer(value: object) -> str:
    return f"{int(round(float(value))):,}".replace(",", ".")


def _format_percentage(value: object) -> str:
    if pd.isna(value):
        return "—"
    return f"{float(value):.1f}%".replace(".", ",")


def _draw_pdf_footer(canvas, document) -> None:
    from reportlab.lib import colors
    from reportlab.lib.units import mm

    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#607784"))
    canvas.drawRightString(
        document.pagesize[0] - document.rightMargin,
        7 * mm,
        f"Página {document.page}",
    )
    canvas.restoreState()
