"""Gera a visão mensal de postos de trabalho em Excel e PDF.

Regra mensal confirmada:

* efetivo mensal = soma dos efetivos semanais / quantidade de semanas;
* trabalhado mensal = soma dos trabalhados semanais / quantidade de semanas;
* ausência mensal = soma das ausências semanais;
* absenteísmo = soma das ausências / soma dos efetivos semanais * 100.

O número de semanas considera semanas distintas depois da aplicação dos
filtros, nunca a quantidade de linhas/oficinas do DataFrame.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
import math
import re
from typing import Final

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from app_postos.core.config import Columns
from app_common.formatting import safe_div
from app_postos.services.export_service import DataExportError


class MonthlyReportColumns:
    ANO = "ano"
    ANO_MES = "ano_mes"
    MES = "mes"
    SEMANAS = "semanas"
    EFETIVO_MEDIO = "efetivo_medio"
    TRABALHADO_MEDIO = "trabalhado_medio"
    TOTAL_AUSENCIA = "total_ausencia"
    ABSENTEISMO = "absenteismo"
    SOMA_EFETIVOS = "soma_efetivos"
    SOMA_TRABALHADOS = "soma_trabalhados"


_REQUIRED_COLUMNS: Final[tuple[str, ...]] = (
    Columns.ANO,
    Columns.ANO_MES,
    Columns.SEMANA,
    Columns.QTD_EFETIVOS,
    Columns.QTD_TRABALHADOS,
)
_VISIBLE_COLUMNS: Final[tuple[tuple[str, str], ...]] = (
    (MonthlyReportColumns.ANO, "Ano"),
    (MonthlyReportColumns.MES, "Mês"),
    (MonthlyReportColumns.SEMANAS, "Semanas (N)"),
    (MonthlyReportColumns.EFETIVO_MEDIO, "Efetivo Médio"),
    (MonthlyReportColumns.TRABALHADO_MEDIO, "Trabalhado Médio"),
    (MonthlyReportColumns.TOTAL_AUSENCIA, "Total Ausência"),
    (MonthlyReportColumns.ABSENTEISMO, "Absenteísmo"),
)
_MONTH_NAMES: Final[dict[int, str]] = {
    1: "Janeiro",
    2: "Fevereiro",
    3: "Março",
    4: "Abril",
    5: "Maio",
    6: "Junho",
    7: "Julho",
    8: "Agosto",
    9: "Setembro",
    10: "Outubro",
    11: "Novembro",
    12: "Dezembro",
}
_MONTH_PATTERN: Final[re.Pattern[str]] = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])$")
_HEADER_FILL: Final[str] = "0E2A3F"
_HEADER_FONT: Final[str] = "FFFFFF"


def build_monthly_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Calcula a visão mensal conforme a regra de negócio confirmada."""
    calculated = _build_monthly_calculation(df)
    return calculated[[column for column, _ in _VISIBLE_COLUMNS]].copy()


def build_monthly_excel_report(df: pd.DataFrame, scope_description: str) -> bytes:
    """Gera o relatório mensal em XLSX com fórmulas auditáveis."""
    calculated = _build_monthly_calculation(df)
    try:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Resumo Mensal"
        _write_excel_summary(worksheet, calculated)
        _write_excel_info(workbook, scope_description, len(calculated))
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    except DataExportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise DataExportError(
            f"Não foi possível montar o resumo mensal em Excel: {exc}"
        ) from exc


def build_monthly_pdf_report(df: pd.DataFrame, scope_description: str) -> bytes:
    """Gera o relatório mensal em PDF com cabeçalho repetido entre páginas."""
    summary = build_monthly_summary(df)
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    except ImportError as exc:
        raise DataExportError(
            "A geração de PDF não está disponível porque a dependência reportlab não foi instalada."
        ) from exc

    try:
        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=11 * mm,
            leftMargin=11 * mm,
            topMargin=13 * mm,
            bottomMargin=13 * mm,
            title="Resumo mensal de postos de trabalho",
            author="Gestão de Postos de Trabalho",
            subject=str(scope_description),
        )

        table_data: list[list[object]] = [[label for _, label in _VISIBLE_COLUMNS]]
        for row in summary.itertuples(index=False):
            table_data.append(
                [
                    str(int(row.ano)),
                    row.mes,
                    str(int(row.semanas)),
                    _format_decimal(row.efetivo_medio),
                    _format_decimal(row.trabalhado_medio),
                    _format_integer(row.total_ausencia),
                    _format_percentage(row.absenteismo),
                ]
            )

        table = Table(
            table_data,
            repeatRows=1,
            colWidths=[19 * mm, 32 * mm, 25 * mm, 35 * mm, 39 * mm, 34 * mm, 31 * mm],
            hAlign="CENTER",
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_HEADER_FILL}")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 8.5),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("FONTSIZE", (0, 1), (-1, -1), 9),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{_HEADER_FILL}")),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )

        story = [table]
        document.build(story, onFirstPage=_draw_pdf_footer, onLaterPages=_draw_pdf_footer)
        return output.getvalue()
    except DataExportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise DataExportError(
            f"Não foi possível montar o resumo mensal em PDF: {exc}"
        ) from exc


def _build_monthly_calculation(df: pd.DataFrame) -> pd.DataFrame:
    source = _prepare_source_data(df)

    weekly = (
        source.groupby(
            [Columns.ANO, Columns.ANO_MES, Columns.SEMANA],
            as_index=False,
        )[[Columns.QTD_EFETIVOS, Columns.QTD_TRABALHADOS]]
        .sum()
        .sort_values([Columns.ANO, Columns.ANO_MES, Columns.SEMANA], kind="stable")
    )
    calculated = (
        weekly.groupby([Columns.ANO, Columns.ANO_MES], as_index=False)
        .agg(
            semanas=(Columns.SEMANA, "nunique"),
            soma_efetivos=(Columns.QTD_EFETIVOS, "sum"),
            soma_trabalhados=(Columns.QTD_TRABALHADOS, "sum"),
        )
        .sort_values([Columns.ANO, Columns.ANO_MES], kind="stable")
        .reset_index(drop=True)
        .rename(
            columns={
                Columns.ANO: MonthlyReportColumns.ANO,
                Columns.ANO_MES: MonthlyReportColumns.ANO_MES,
            }
        )
    )
    calculated[MonthlyReportColumns.MES] = calculated[MonthlyReportColumns.ANO_MES].map(
        _format_month_label
    )
    calculated[MonthlyReportColumns.EFETIVO_MEDIO] = (
        calculated[MonthlyReportColumns.SOMA_EFETIVOS]
        / calculated[MonthlyReportColumns.SEMANAS]
    )
    calculated[MonthlyReportColumns.TRABALHADO_MEDIO] = (
        calculated[MonthlyReportColumns.SOMA_TRABALHADOS]
        / calculated[MonthlyReportColumns.SEMANAS]
    )
    calculated[MonthlyReportColumns.TOTAL_AUSENCIA] = (
        calculated[MonthlyReportColumns.SOMA_EFETIVOS]
        - calculated[MonthlyReportColumns.SOMA_TRABALHADOS]
    )
    calculated[MonthlyReportColumns.ABSENTEISMO] = calculated.apply(
        lambda row: safe_div(
            row[MonthlyReportColumns.TOTAL_AUSENCIA],
            row[MonthlyReportColumns.SOMA_EFETIVOS],
        )
        * 100,
        axis=1,
    )
    return calculated[
        [
            MonthlyReportColumns.ANO,
            MonthlyReportColumns.ANO_MES,
            MonthlyReportColumns.MES,
            MonthlyReportColumns.SEMANAS,
            MonthlyReportColumns.EFETIVO_MEDIO,
            MonthlyReportColumns.TRABALHADO_MEDIO,
            MonthlyReportColumns.TOTAL_AUSENCIA,
            MonthlyReportColumns.ABSENTEISMO,
            MonthlyReportColumns.SOMA_EFETIVOS,
            MonthlyReportColumns.SOMA_TRABALHADOS,
        ]
    ]


def _prepare_source_data(df: pd.DataFrame) -> pd.DataFrame:
    if not isinstance(df, pd.DataFrame):
        raise DataExportError("Os dados informados para o relatório mensal são inválidos.")
    if df.empty:
        raise DataExportError("Não há registros para exportar.")

    missing = [column for column in _REQUIRED_COLUMNS if column not in df.columns]
    if missing:
        raise DataExportError(
            "Os dados não possuem todas as colunas necessárias para o resumo mensal: "
            f"{', '.join(missing)}."
        )

    numeric_columns = [
        Columns.ANO,
        Columns.SEMANA,
        Columns.QTD_EFETIVOS,
        Columns.QTD_TRABALHADOS,
    ]
    converted = {
        column: pd.to_numeric(df[column], errors="coerce")
        for column in numeric_columns
    }
    invalid = [column for column, values in converted.items() if values.isna().any()]
    if invalid:
        raise DataExportError(
            "Há valores vazios ou não numéricos nas colunas do resumo mensal: "
            f"{', '.join(invalid)}."
        )
    non_finite = [
        column
        for column, values in converted.items()
        if (~values.map(lambda value: math.isfinite(float(value)))).any()
    ]
    if non_finite:
        raise DataExportError(
            "Há valores infinitos nas colunas do resumo mensal: "
            f"{', '.join(non_finite)}."
        )

    if ((converted[Columns.SEMANA] % 1 != 0) | ~converted[Columns.SEMANA].between(1, 53)).any():
        raise DataExportError("A coluna semana deve conter números inteiros entre 1 e 53.")
    if ((converted[Columns.ANO] % 1 != 0) | ~converted[Columns.ANO].between(1, 9999)).any():
        raise DataExportError("A coluna ano deve conter anos válidos.")
    for column in (Columns.QTD_EFETIVOS, Columns.QTD_TRABALHADOS):
        if (converted[column] < 0).any():
            raise DataExportError(f"A coluna {column} não pode conter valores negativos.")
        if (converted[column] % 1 != 0).any():
            raise DataExportError(f"A coluna {column} deve conter quantidades inteiras.")

    periods = df[Columns.ANO_MES].astype("string").str.strip()
    if periods.isna().any() or (~periods.str.match(_MONTH_PATTERN)).any():
        raise DataExportError("A coluna ano_mes deve usar o formato válido AAAA-MM.")
    period_years = periods.str.slice(0, 4).astype("int64")
    normalized_years = converted[Columns.ANO].astype("int64")
    if (period_years != normalized_years).any():
        raise DataExportError("O ano de ano_mes não corresponde à coluna ano.")

    normalized = pd.DataFrame(converted, index=df.index)
    for column in numeric_columns:
        normalized[column] = normalized[column].astype("int64")
    normalized[Columns.ANO_MES] = periods

    month_count_by_week = normalized.groupby(
        [Columns.ANO, Columns.SEMANA]
    )[Columns.ANO_MES].nunique()
    if (month_count_by_week > 1).any():
        raise DataExportError(
            "Uma mesma semana está associada a mais de um mês no recorte selecionado."
        )
    return normalized


def _write_excel_summary(worksheet, calculated: pd.DataFrame) -> None:
    headers = [label for _, label in _VISIBLE_COLUMNS]
    worksheet.append([*headers, "Soma Efetivos", "Soma Trabalhados"])

    for row_index, row in enumerate(calculated.itertuples(index=False), start=2):
        worksheet.append(
            [
                int(row.ano),
                row.mes,
                int(row.semanas),
                f"=H{row_index}/C{row_index}",
                f"=I{row_index}/C{row_index}",
                f"=H{row_index}-I{row_index}",
                f'=IF(H{row_index}=0,"",F{row_index}/H{row_index})',
                int(row.soma_efetivos),
                int(row.soma_trabalhados),
            ]
        )

    header_fill = PatternFill(fill_type="solid", fgColor=_HEADER_FILL)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color=_HEADER_FONT, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    border = Border(
        left=Side(style="thin", color=_HEADER_FILL),
        right=Side(style="thin", color=_HEADER_FILL),
        top=Side(style="thin", color=_HEADER_FILL),
        bottom=Side(style="thin", color=_HEADER_FILL),
    )
    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = border
            if cell.row > 1:
                cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFF")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:G{worksheet.max_row}"
    worksheet.row_dimensions[1].height = 24
    widths = {"A": 10, "B": 22, "C": 15, "D": 19, "E": 21, "F": 18, "G": 16}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    for column in ("D", "E"):
        for cell in worksheet[column][1:]:
            cell.number_format = "#,##0.0"
    for cell in worksheet["F"][1:]:
        cell.number_format = "#,##0"
    for cell in worksheet["G"][1:]:
        cell.number_format = "0.0%"
    for row in worksheet.iter_rows(min_row=2, max_col=7):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.column_dimensions["H"].hidden = True
    worksheet.column_dimensions["I"].hidden = True
    table = Table(displayName="ResumoMensal", ref=f"A1:G{worksheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _write_excel_info(workbook: Workbook, scope_description: str, row_count: int) -> None:
    worksheet = workbook.create_sheet("Informações")
    worksheet.append(["Relatório mensal"])
    worksheet.append(["Escopo", str(scope_description)])
    worksheet.append(["Meses exportados", int(row_count)])
    worksheet.append(["Gerado em", datetime.now().astimezone().strftime("%d/%m/%Y %H:%M:%S %z")])
    worksheet.append(
        [
            "Regra",
            "Efetivo e trabalhado = soma semanal / N; ausência = soma semanal; "
            "absenteísmo = soma das ausências / soma dos efetivos semanais.",
        ]
    )
    worksheet.append(["N", "Quantidade de semanas distintas no mês após os filtros."])
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 92
    worksheet["A1"].font = Font(bold=True, size=14, color=_HEADER_FONT)
    worksheet["A1"].fill = PatternFill(fill_type="solid", fgColor=_HEADER_FILL)
    worksheet.merge_cells("A1:B1")


def _format_month_label(period: str) -> str:
    year, month = str(period).split("-")
    return f"{_MONTH_NAMES[int(month)]}/{year}"


def _format_decimal(value: object) -> str:
    return f"{float(value):,.1f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _format_integer(value: object) -> str:
    return f"{int(round(float(value))):,}".replace(",", ".")


def _format_percentage(value: object) -> str:
    if pd.isna(value):
        return "—"
    return f"{float(value):.1f}%".replace(".", ",")


def _draw_pdf_footer(canvas, document) -> None:
    from reportlab.lib import colors
    from reportlab.lib.units import mm

    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#607784"))
    canvas.drawRightString(
        document.pagesize[0] - document.rightMargin,
        7 * mm,
        f"Página {document.page}",
    )
    canvas.restoreState()
