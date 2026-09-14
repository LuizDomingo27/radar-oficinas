from __future__ import annotations

from io import BytesIO
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from app_postos.core.config import Columns
from app_postos.services.export_service import DataExportError
from app_postos.services.monthly_report_service import (
    MonthlyReportColumns,
    build_monthly_excel_report,
    build_monthly_pdf_report,
    build_monthly_summary,
)


def _monthly_dataframe() -> pd.DataFrame:
    return pd.DataFrame(
        {
            Columns.ANO: [2026, 2026, 2026, 2026, 2026, 2026],
            Columns.ANO_MES: ["2026-01", "2026-01", "2026-01", "2026-01", "2026-02", "2026-02"],
            Columns.SEMANA: [1, 2, 3, 4, 5, 5],
            Columns.QTD_EFETIVOS: [250, 260, 255, 265, 100, 50],
            Columns.QTD_TRABALHADOS: [235, 240, 238, 242, 90, 45],
        }
    )


class MonthlySummaryTests(unittest.TestCase):
    def test_applies_confirmed_monthly_rule(self) -> None:
        result = build_monthly_summary(_monthly_dataframe())

        january = result[result[MonthlyReportColumns.MES] == "Janeiro/2026"].iloc[0]
        self.assertEqual(january[MonthlyReportColumns.SEMANAS], 4)
        self.assertAlmostEqual(january[MonthlyReportColumns.EFETIVO_MEDIO], 257.5)
        self.assertAlmostEqual(january[MonthlyReportColumns.TRABALHADO_MEDIO], 238.75)
        self.assertEqual(january[MonthlyReportColumns.TOTAL_AUSENCIA], 75)
        self.assertAlmostEqual(january[MonthlyReportColumns.ABSENTEISMO], 75 / 1030 * 100)

    def test_counts_distinct_weeks_instead_of_rows(self) -> None:
        result = build_monthly_summary(_monthly_dataframe())

        february = result[result[MonthlyReportColumns.MES] == "Fevereiro/2026"].iloc[0]
        self.assertEqual(february[MonthlyReportColumns.SEMANAS], 1)
        self.assertEqual(february[MonthlyReportColumns.EFETIVO_MEDIO], 150)
        self.assertEqual(february[MonthlyReportColumns.TRABALHADO_MEDIO], 135)
        self.assertEqual(february[MonthlyReportColumns.TOTAL_AUSENCIA], 15)
        self.assertAlmostEqual(february[MonthlyReportColumns.ABSENTEISMO], 10)

    def test_zero_effective_produces_missing_absenteeism(self) -> None:
        data = _monthly_dataframe().iloc[[0]].copy()
        data[Columns.QTD_EFETIVOS] = 0
        data[Columns.QTD_TRABALHADOS] = 0

        result = build_monthly_summary(data)

        self.assertTrue(pd.isna(result.iloc[0][MonthlyReportColumns.ABSENTEISMO]))


class MonthlyValidationTests(unittest.TestCase):
    def test_rejects_invalid_period(self) -> None:
        data = _monthly_dataframe()
        data.loc[0, Columns.ANO_MES] = "01/2026"
        with self.assertRaisesRegex(DataExportError, "AAAA-MM"):
            build_monthly_summary(data)

    def test_rejects_period_from_another_year(self) -> None:
        data = _monthly_dataframe()
        data.loc[0, Columns.ANO_MES] = "2025-01"
        with self.assertRaisesRegex(DataExportError, "não corresponde"):
            build_monthly_summary(data)

    def test_rejects_same_week_assigned_to_two_months(self) -> None:
        data = _monthly_dataframe()
        duplicate = data.iloc[[0]].copy()
        duplicate[Columns.ANO_MES] = "2026-02"
        data = pd.concat([data, duplicate], ignore_index=True)
        with self.assertRaisesRegex(DataExportError, "mesma semana"):
            build_monthly_summary(data)

    def test_rejects_missing_columns(self) -> None:
        with self.assertRaisesRegex(DataExportError, "colunas necessárias"):
            build_monthly_summary(pd.DataFrame({Columns.ANO: [2026]}))


class MonthlyExcelTests(unittest.TestCase):
    def test_builds_formula_driven_workbook(self) -> None:
        content = build_monthly_excel_report(_monthly_dataframe(), "Ano 2026")
        workbook = load_workbook(BytesIO(content), data_only=False)
        worksheet = workbook["Resumo Mensal"]

        self.assertEqual(
            [cell.value for cell in worksheet[1]][:7],
            [
                "Ano",
                "Mês",
                "Semanas (N)",
                "Efetivo Médio",
                "Trabalhado Médio",
                "Total Ausência",
                "Absenteísmo",
            ],
        )
        self.assertEqual(worksheet["D2"].value, "=H2/C2")
        self.assertEqual(worksheet["E2"].value, "=I2/C2")
        self.assertEqual(worksheet["F2"].value, "=H2-I2")
        self.assertEqual(worksheet["G2"].value, '=IF(H2=0,"",F2/H2)')
        self.assertTrue(worksheet.column_dimensions["H"].hidden)
        self.assertTrue(worksheet.column_dimensions["I"].hidden)
        self.assertEqual(worksheet.tables["ResumoMensal"].ref, "A1:G3")
        self.assertEqual(worksheet["A1"].fill.fgColor.rgb, "000E2A3F")
        self.assertEqual(worksheet["A1"].font.color.rgb, "00FFFFFF")
        self.assertEqual(worksheet["A2"].fill.fgColor.rgb, "00FFFFFF")
        self.assertEqual(worksheet["A2"].border.left.style, "thin")
        self.assertFalse(worksheet.tables["ResumoMensal"].tableStyleInfo.showRowStripes)

    def test_wraps_excel_failure(self) -> None:
        with patch("app_postos.services.monthly_report_service.Workbook.save", side_effect=OSError("disco")):
            with self.assertRaisesRegex(DataExportError, "resumo mensal em Excel"):
                build_monthly_excel_report(_monthly_dataframe(), "Ano 2026")


class MonthlyPdfTests(unittest.TestCase):
    def test_builds_pdf_with_expected_monthly_values(self) -> None:
        try:
            from pypdf import PdfReader
        except ImportError:  # pragma: no cover
            self.skipTest("pypdf não disponível para validar o PDF")

        content = build_monthly_pdf_report(_monthly_dataframe(), "Ano 2026")
        text = "\n".join(
            page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages
        )

        self.assertTrue(content.startswith(b"%PDF"))
        self.assertIn("Janeiro/2026", text)
        self.assertIn("257,5", text)
        self.assertIn("238,8", text)
        self.assertIn("75", text)
        self.assertIn("7,3%", text)

    def test_wraps_pdf_failure(self) -> None:
        with patch(
            "reportlab.platypus.SimpleDocTemplate.build",
            side_effect=OSError("render"),
        ):
            with self.assertRaisesRegex(DataExportError, "resumo mensal em PDF"):
                build_monthly_pdf_report(_monthly_dataframe(), "Ano 2026")


if __name__ == "__main__":
    unittest.main()
