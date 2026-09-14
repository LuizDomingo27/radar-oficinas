from __future__ import annotations

from io import BytesIO
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from app_postos.core.config import Columns
from app_postos.services.export_service import DataExportError
from app_postos.services.weekly_report_service import (
    WeeklyReportColumns,
    build_weekly_excel_report,
    build_weekly_pdf_report,
    build_weekly_summary,
)


def _valid_dataframe() -> pd.DataFrame:
    return pd.DataFrame(
        {
            Columns.ANO: [2026, 2026, 2026, 2025],
            Columns.SEMANA: [30, 30, 31, 30],
            Columns.QTD_EFETIVOS: [100, 156, 0, 50],
            Columns.QTD_TRABALHADOS: [80, 120, 0, 45],
        }
    )


class WeeklySummaryTests(unittest.TestCase):
    def test_groups_by_year_and_week_and_calculates_metrics(self) -> None:
        source = _valid_dataframe()
        original = source.copy(deep=True)

        result = build_weekly_summary(source)

        self.assertEqual(len(result), 3)
        row = result[(result["ano"] == 2026) & (result["semana"] == 30)].iloc[0]
        self.assertEqual(row[WeeklyReportColumns.TOTAL_EFETIVO], 256)
        self.assertEqual(row[WeeklyReportColumns.TOTAL_TRABALHADO], 200)
        self.assertEqual(row[WeeklyReportColumns.TOTAL_AUSENCIA], 56)
        self.assertAlmostEqual(row[WeeklyReportColumns.ABSENTEISMO], 21.875)
        pd.testing.assert_frame_equal(source, original)

    def test_zero_effective_produces_missing_absenteeism_without_crashing(self) -> None:
        result = build_weekly_summary(_valid_dataframe())

        row = result[(result["ano"] == 2026) & (result["semana"] == 31)].iloc[0]
        self.assertTrue(pd.isna(row[WeeklyReportColumns.ABSENTEISMO]))

    def test_normalizes_numeric_text_before_aggregation(self) -> None:
        data = _valid_dataframe().astype(str)

        result = build_weekly_summary(data)

        row = result[(result["ano"] == 2026) & (result["semana"] == 30)].iloc[0]
        self.assertEqual(row[WeeklyReportColumns.TOTAL_EFETIVO], 256)
        self.assertEqual(row[WeeklyReportColumns.TOTAL_TRABALHADO], 200)

    def test_does_not_merge_the_same_week_from_different_years(self) -> None:
        result = build_weekly_summary(_valid_dataframe())

        week_30 = result[result[WeeklyReportColumns.SEMANA] == 30]
        self.assertEqual(week_30[WeeklyReportColumns.ANO].tolist(), [2025, 2026])


class WeeklyReportValidationTests(unittest.TestCase):
    def test_rejects_non_dataframe(self) -> None:
        with self.assertRaisesRegex(DataExportError, "dados informados"):
            build_weekly_summary([])  # type: ignore[arg-type]

    def test_rejects_empty_dataframe(self) -> None:
        with self.assertRaisesRegex(DataExportError, "Não há registros"):
            build_weekly_summary(pd.DataFrame())

    def test_reports_all_missing_columns(self) -> None:
        with self.assertRaises(DataExportError) as raised:
            build_weekly_summary(pd.DataFrame({Columns.ANO: [2026]}))

        message = str(raised.exception)
        self.assertIn(Columns.SEMANA, message)
        self.assertIn(Columns.QTD_EFETIVOS, message)
        self.assertIn(Columns.QTD_TRABALHADOS, message)

    def test_rejects_invalid_week(self) -> None:
        data = _valid_dataframe()
        data.loc[0, Columns.SEMANA] = 54
        with self.assertRaisesRegex(DataExportError, "entre 1 e 53"):
            build_weekly_summary(data)

    def test_rejects_fractional_week(self) -> None:
        data = _valid_dataframe()
        data[Columns.SEMANA] = data[Columns.SEMANA].astype("float64")
        data.loc[0, Columns.SEMANA] = 30.5
        with self.assertRaisesRegex(DataExportError, "inteiros"):
            build_weekly_summary(data)

    def test_rejects_non_numeric_value(self) -> None:
        data = _valid_dataframe()
        data[Columns.QTD_EFETIVOS] = data[Columns.QTD_EFETIVOS].astype("object")
        data.loc[0, Columns.QTD_EFETIVOS] = "inválido"
        with self.assertRaisesRegex(DataExportError, "não numéricos"):
            build_weekly_summary(data)

    def test_rejects_negative_quantities(self) -> None:
        data = _valid_dataframe()
        data.loc[0, Columns.QTD_TRABALHADOS] = -1
        with self.assertRaisesRegex(DataExportError, "não pode conter valores negativos"):
            build_weekly_summary(data)

    def test_rejects_infinite_values(self) -> None:
        data = _valid_dataframe()
        data[Columns.QTD_EFETIVOS] = data[Columns.QTD_EFETIVOS].astype("float64")
        data.loc[0, Columns.QTD_EFETIVOS] = float("inf")
        with self.assertRaisesRegex(DataExportError, "infinitos"):
            build_weekly_summary(data)

    def test_rejects_fractional_quantities(self) -> None:
        data = _valid_dataframe()
        data[Columns.QTD_TRABALHADOS] = data[Columns.QTD_TRABALHADOS].astype("float64")
        data.loc[0, Columns.QTD_TRABALHADOS] = 80.5
        with self.assertRaisesRegex(DataExportError, "quantidades inteiras"):
            build_weekly_summary(data)


class WeeklyExcelReportTests(unittest.TestCase):
    def test_builds_workbook_with_formulas_formats_and_metadata(self) -> None:
        content = build_weekly_excel_report(_valid_dataframe(), "Filtros atuais")
        workbook = load_workbook(BytesIO(content), data_only=False)
        worksheet = workbook["Resumo Semanal"]

        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            ["Ano", "Semana", "Total Efetivo", "Total Trabalhado", "Total Ausência", "Absenteísmo"],
        )
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.tables["ResumoSemanal"].ref, "A1:F4")
        self.assertEqual(worksheet["E3"].value, "=C3-D3")
        self.assertEqual(worksheet["F3"].value, '=IF(C3=0,"",E3/C3)')
        self.assertEqual(worksheet["F3"].number_format, "0.0%")
        self.assertEqual(worksheet["A1"].fill.fgColor.rgb, "000E2A3F")
        self.assertEqual(worksheet["A1"].font.color.rgb, "00FFFFFF")
        self.assertEqual(worksheet["A2"].fill.fgColor.rgb, "00FFFFFF")
        self.assertEqual(worksheet["A2"].border.left.style, "thin")
        self.assertFalse(worksheet.tables["ResumoSemanal"].tableStyleInfo.showRowStripes)
        self.assertEqual(workbook["Informações"]["B2"].value, "Filtros atuais")
        self.assertEqual(workbook.calculation.calcMode, "auto")
        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        self.assertGreater(len(content), 1000)

    def test_wraps_workbook_failures_in_domain_error(self) -> None:
        with patch("app_postos.services.weekly_report_service.Workbook.save", side_effect=OSError("disco")):
            with self.assertRaisesRegex(DataExportError, "resumo semanal em Excel"):
                build_weekly_excel_report(_valid_dataframe(), "Filtros atuais")


class WeeklyPdfReportTests(unittest.TestCase):
    def test_builds_readable_pdf_with_expected_values(self) -> None:
        try:
            from pypdf import PdfReader
        except ImportError:  # pragma: no cover - dependência presente no ambiente de CI do projeto.
            self.skipTest("pypdf não disponível para validar o PDF")

        content = build_weekly_pdf_report(_valid_dataframe(), "Filtros atuais")
        reader = PdfReader(BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)

        self.assertTrue(content.startswith(b"%PDF"))
        self.assertIn("Semana", text)
        self.assertIn("256", text)
        self.assertIn("200", text)
        self.assertIn("56", text)
        self.assertIn("21,9%", text)

    def test_wraps_pdf_build_failures_in_domain_error(self) -> None:
        with patch(
            "reportlab.platypus.SimpleDocTemplate.build",
            side_effect=OSError("falha de renderização"),
        ):
            with self.assertRaisesRegex(DataExportError, "resumo semanal em PDF"):
                build_weekly_pdf_report(_valid_dataframe(), "Filtros atuais")

    def test_repeats_header_when_pdf_spans_multiple_pages(self) -> None:
        try:
            from pypdf import PdfReader
        except ImportError:  # pragma: no cover
            self.skipTest("pypdf não disponível para validar o PDF")

        rows = [
            {
                Columns.ANO: 2026,
                Columns.SEMANA: (index % 53) + 1,
                Columns.QTD_EFETIVOS: 100 + index,
                Columns.QTD_TRABALHADOS: 90 + index,
            }
            for index in range(53)
        ]
        content = build_weekly_pdf_report(pd.DataFrame(rows), "Ano 2026")
        pages = PdfReader(BytesIO(content)).pages

        self.assertGreater(len(pages), 1)
        for page in pages:
            text = page.extract_text() or ""
            self.assertIn("Semana", text)
            self.assertIn("Absenteísmo", text)


if __name__ == "__main__":
    unittest.main()
